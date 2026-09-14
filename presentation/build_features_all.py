"""
Generates features_all.pptx -- a comprehensive, feature-by-feature walkthrough
of every phase of the Telecom Churn / Retention Intelligence system, current
as of the actual repository state (Phases 1-8 all implemented).

Run:
    python build_features_all.py
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

OUT_PATH = Path(__file__).resolve().parent / "features_all.pptx"

# ---------------------------------------------------------------- palette --
STAGE = RGBColor(0x0B, 0x0F, 0x14)
SURFACE = RGBColor(0x12, 0x1A, 0x22)
SURFACE_2 = RGBColor(0x18, 0x22, 0x2C)
LINE = RGBColor(0x24, 0x32, 0x3E)
INK = RGBColor(0xEE, 0xF3, 0xF5)
INK_SOFT = RGBColor(0xAA, 0xBA, 0xC2)
INK_FAINT = RGBColor(0x6C, 0x7E, 0x87)
ACCENT = RGBColor(0x33, 0xD6, 0xC6)
ACCENT_STRONG = RGBColor(0x94, 0xF5, 0xEA)
ENERGY = RGBColor(0xFF, 0x8A, 0x5C)
GOOD = RGBColor(0x5E, 0xD6, 0x8A)
WARN = RGBColor(0xF2, 0xC2, 0x4B)

HEAD_FONT = "Segoe UI Semibold"
BODY_FONT = "Segoe UI"
MONO_FONT = "Consolas"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

STATUS_COLOR = {"done": GOOD, "warn": WARN, "new": ACCENT}
STATUS_LABEL = {"done": "BUILT", "warn": "BUILT · CAVEAT", "new": "NEW · PHASE 7/8"}


def new_deck() -> Presentation:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def add_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = STAGE
    bg.line.fill.background()
    bg.shadow.inherit = False
    slide.shapes._spTree.remove(bg._element)
    slide.shapes._spTree.insert(2, bg._element)
    return slide


def add_text(slide, text, left, top, width, height, size, color=INK,
             bold=False, italic=False, font=BODY_FONT, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, line_spacing=1.15):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font
    run.font.color.rgb = color
    return box


def eyebrow(slide, text, top=Inches(0.55)):
    return add_text(slide, text.upper(), Inches(0.9), top, SLIDE_W - Inches(1.8), Inches(0.4),
                     13, ACCENT, bold=True, font=MONO_FONT)


def hr(slide, left, top, width, color=LINE):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, left, top, left + width, top)
    ln.line.color.rgb = color
    ln.line.width = Pt(1)
    return ln


def status_pill(slide, status, right_edge=SLIDE_W - Inches(0.9), top=Inches(0.5)):
    label = STATUS_LABEL[status]
    color = STATUS_COLOR[status]
    w = Inches(0.14 * len(label) + 0.35)
    left = right_edge - w
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, Inches(0.34))
    pill.adjustments[0] = 0.5
    pill.fill.solid()
    pill.fill.fore_color.rgb = SURFACE_2
    pill.line.color.rgb = color
    pill.line.width = Pt(1)
    pill.shadow.inherit = False
    tf = pill.text_frame
    tf.word_wrap = False
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.name = MONO_FONT
    r.font.color.rgb = color
    return pill


def section_header(slide, kicker, title, status=None):
    eyebrow(slide, kicker)
    if status:
        status_pill(slide, status)
    add_text(slide, title, Inches(0.9), Inches(0.92), SLIDE_W - Inches(1.8), Inches(0.75),
              30, INK, bold=True, font=HEAD_FONT)
    hr(slide, Inches(0.9), Inches(1.62), SLIDE_W - Inches(1.8))


def bullet_block(slide, left, top, width, height, items, size=14.5, gap=Pt(6)):
    """items: list of (bold_lead, rest) or plain strings."""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = gap
        p.line_spacing = 1.08
        dot = p.add_run()
        dot.text = "▸  "
        dot.font.size = Pt(size)
        dot.font.color.rgb = ACCENT
        dot.font.bold = True
        if isinstance(item, tuple):
            lead, rest = item
            r1 = p.add_run()
            r1.text = lead
            r1.font.size = Pt(size)
            r1.font.bold = True
            r1.font.color.rgb = INK
            r1.font.name = BODY_FONT
            r2 = p.add_run()
            r2.text = rest
            r2.font.size = Pt(size)
            r2.font.color.rgb = INK_SOFT
            r2.font.name = BODY_FONT
        else:
            r1 = p.add_run()
            r1.text = item
            r1.font.size = Pt(size)
            r1.font.color.rgb = INK_SOFT
            r1.font.name = BODY_FONT
    return box


def code_card(slide, left, top, width, height, lines, title=None):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.adjustments[0] = 0.04
    card.fill.solid()
    card.fill.fore_color.rgb = SURFACE
    card.line.color.rgb = LINE
    card.line.width = Pt(1)
    card.shadow.inherit = False
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.22)
    tf.margin_right = Inches(0.18)
    tf.margin_top = Inches(0.16)
    for i, line in enumerate(([title] if title else []) + [""] * (1 if title else 0) + lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.15
        r = p.add_run()
        r.text = line
        r.font.name = MONO_FONT
        r.font.size = Pt(12.5)
        if title and i == 0:
            r.font.color.rgb = ACCENT_STRONG
            r.font.bold = True
        else:
            r.font.color.rgb = INK_SOFT
    return card


def table_slide(prs, kicker, title, status, headers, rows, col_widths, note=None):
    s = add_slide(prs)
    section_header(s, kicker, title, status)
    top = Inches(1.95)
    left = Inches(0.9)
    total_w = sum(col_widths, Inches(0))
    n_rows = len(rows) + 1
    tbl_h = Inches(0.5) * n_rows
    tbl_h = min(tbl_h, Inches(4.9))
    gfx = s.shapes.add_table(n_rows, len(headers), left, top, total_w, tbl_h)
    table = gfx.table
    for i, w in enumerate(col_widths):
        table.columns[i].width = w
    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = SURFACE_2
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        run = p.runs[0]
        run.font.bold = True
        run.font.size = Pt(12.5)
        run.font.color.rgb = ACCENT
        run.font.name = MONO_FONT
        cell.margin_top = Inches(0.05)
        cell.margin_bottom = Inches(0.05)
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(val)
            cell.fill.solid()
            cell.fill.fore_color.rgb = STAGE if r % 2 else SURFACE
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            run = p.runs[0]
            run.font.size = Pt(12)
            run.font.color.rgb = INK_SOFT if c else INK
            run.font.name = BODY_FONT
            cell.margin_top = Inches(0.04)
            cell.margin_bottom = Inches(0.04)
    if note:
        add_text(s, note, Inches(0.9), top + tbl_h + Inches(0.25), SLIDE_W - Inches(1.8), Inches(0.8),
                  12.5, INK_FAINT, italic=True, font=BODY_FONT)
    return s


def bullets_slide(prs, kicker, title, status, items, code=None, code_title=None):
    s = add_slide(prs)
    section_header(s, kicker, title, status)
    text_w = Inches(6.6) if code else SLIDE_W - Inches(1.8)
    bullet_block(s, Inches(0.9), Inches(2.0), text_w, Inches(4.8), items)
    if code:
        code_card(s, Inches(7.75), Inches(2.0), SLIDE_W - Inches(7.75) - Inches(0.9), Inches(4.6),
                  code, title=code_title)
    return s


def build() -> None:
    prs = new_deck()

    # ---------------------------------------------------------- 1. TITLE --
    s = add_slide(prs)
    add_text(s, "RETENTION INTELLIGENCE SYSTEM", Inches(1), Inches(2.15), SLIDE_W - Inches(2), Inches(0.5),
              15, ACCENT, bold=True, font=MONO_FONT, align=PP_ALIGN.CENTER)
    add_text(s, "All Features, Explained", Inches(1), Inches(2.65), SLIDE_W - Inches(2), Inches(1.3),
              50, INK, bold=True, font=HEAD_FONT, align=PP_ALIGN.CENTER)
    add_text(s, "From a raw CSV to a governed, tool-calling AI assistant — every phase, "
                 "every file, every guardrail.",
              Inches(2), Inches(3.95), SLIDE_W - Inches(4), Inches(0.7), 17, INK_SOFT,
              font=BODY_FONT, align=PP_ALIGN.CENTER)
    hr(s, Inches(4.5), Inches(4.85), Inches(4.333))
    add_text(s, "IBM / Telco Customer Churn dataset · 7,043 customers · 8 phases · Phases 1–8 all implemented",
              Inches(1), Inches(5.05), SLIDE_W - Inches(2), Inches(0.5), 12.5, INK_FAINT,
              font=MONO_FONT, align=PP_ALIGN.CENTER)

    # ------------------------------------------------------- 2. OVERVIEW --
    s = add_slide(prs)
    section_header(s, "System Map", "One pipeline, eight phases, one API surface")
    phases = [
        ("1", "Python Core", "Clean the raw CSV once, everywhere"),
        ("2", "Database", "Staging → curated MySQL + risk view"),
        ("3", "FastAPI", "Every capability as one callable endpoint"),
        ("4", "React UI", "Support-agent & exec dashboard, 5 tabs"),
        ("5", "Data Eng.", "Repeatable, quality-gated, orchestrated"),
        ("6", "ML Models", "Trained, explainable, batch-scored"),
        ("7", "Claude Audit", "Finds & fixes a real inference bug"),
        ("8", "AI Assistant", "Tool-calling layer, evaluated & hardened"),
    ]
    cols = 4
    card_w, card_h = Inches(2.85), Inches(2.15)
    gap_x, gap_y = Inches(0.2), Inches(0.25)
    left0 = (SLIDE_W - (card_w * cols + gap_x * (cols - 1))) / 2
    top0 = Inches(2.05)
    for i, (num, name, desc) in enumerate(phases):
        col, row = i % cols, i // cols
        left = left0 + col * (card_w + gap_x)
        top = top0 + row * (card_h + gap_y)
        card = slide_card = prs.slides[-1].shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, card_w, card_h)
        card.adjustments[0] = 0.06
        card.fill.solid()
        card.fill.fore_color.rgb = SURFACE
        card.line.color.rgb = ACCENT if row == 1 else LINE
        card.line.width = Pt(1.25 if row == 1 else 1)
        card.shadow.inherit = False
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.16)
        p0 = tf.paragraphs[0]
        r0 = p0.add_run(); r0.text = f"PHASE {num}"
        r0.font.size = Pt(11); r0.font.bold = True; r0.font.name = MONO_FONT; r0.font.color.rgb = ACCENT
        p1 = tf.add_paragraph()
        r1 = p1.add_run(); r1.text = name
        r1.font.size = Pt(17); r1.font.bold = True; r1.font.name = HEAD_FONT; r1.font.color.rgb = INK
        p2 = tf.add_paragraph()
        p2.space_before = Pt(4)
        r2 = p2.add_run(); r2.text = desc
        r2.font.size = Pt(12); r2.font.name = BODY_FONT; r2.font.color.rgb = INK_SOFT

    # ------------------------------------------------- 3. PHASE 1 CORE ----
    bullets_slide(
        prs, "Phase 1 — Python Core", "Turning a messy CRM export into a trustworthy DataFrame", "done",
        items=[
            ("Two real data problems, fixed once: ", "`TotalCharges` arrives as text with 11 blank cells "
             "(tenure = 0 customers not yet billed), and every Yes/No column needs to become 1/0."),
            ("`CustomerCleaner` class: ", "one entry point, `clean()`, runs four steps in order — "
             "`standard()` snake_cases columns, `total_cha()` coerces charges to numeric, `normal()` maps "
             "Yes/No → 1/0, `null_handler()` fills blanks with that row's monthly charge."),
            ("Reused everywhere: ", "imported by the DB loader, the DE pipeline notebooks, and training — "
             "one implementation, never re-derived."),
            ("Canonical pipeline script: ", "`customer_pipeline.py` resolves paths relative to itself, writes "
             "a rotating log, and asserts no nulls in `monthly_charges` and churn is strictly 0/1 before saving."),
        ],
        code=[
            "class CustomerCleaner:",
            "    def __init__(self, df):",
            "        self.df = df.copy()",
            "",
            "    def clean(self):",
            "        self.standard()",
            "        self.total_cha()",
            "        self.normal()",
            "        self.null_handler()",
            "        return self.df",
        ],
        code_title="customer_cleaner.py",
    )

    # ------------------------------------------------- 4. PHASE 2 DB -----
    table_slide(
        prs, "Phase 2 — Database Layer", "A staging → curated MySQL schema, not a flat CSV", "done",
        headers=["Table / View", "Role"],
        rows=[
            ("stg_customer_raw", "Landing zone — exact CSV shape, every column VARCHAR, untouched"),
            ("stg_telco_customer", "Typed, cleaned customer table (ORM class Customer) — GET /customers/{id} reads this"),
            ("ingestion_log", "Audit row per load: source file, row count, LOADED/REJECTED status"),
            ("dim_contract / dim_payment", "Lookup dimensions — one row per distinct value"),
            ("fact_customer_account", "Curated fact table: tenure, charges, churn_flag, joined by surrogate key"),
            ("v_high_risk_customers", "SQL view: month-to-month + tenure<12 + above-average charges (rule-based, not ML)"),
        ],
        col_widths=[Inches(3.4), Inches(7.9)],
        note="SQLAlchemy ORM used two ways: declarative Base/Column models (tables3.py) for SQL3, and every API read goes through a SessionLocal() session rather than a raw connection.",
    )

    # ------------------------------------------------- 5. PHASE 3 API ----
    table_slide(
        prs, "Phase 3 — FastAPI Backend", "One service, every capability as a REST endpoint", "done",
        headers=["Endpoint", "Method", "Purpose"],
        rows=[
            ("/customers/{id}", "GET", "Single customer profile, 404 if unknown"),
            ("/customers", "GET", "Full customer list"),
            ("/customers", "POST \U0001f512", "Create — requires X-API-Key header"),
            ("/customers/{id}", "PATCH \U0001f512", "Partial update — requires X-API-Key"),
            ("/customers/{id}", "DELETE \U0001f512", "Delete — requires X-API-Key"),
            ("/customers/high-risk", "GET", "Reads v_high_risk_customers, filterable by limit/min_tenure/max_tenure"),
            ("/churn/summary", "GET", "Executive KPIs: totals, churn rate, breakdown by contract & internet service"),
            ("/predict-churn", "POST", "Live model risk score"),
            ("/assistant/chat", "POST", "Phase 8: tool-calling Claude assistant turn"),
        ],
        col_widths=[Inches(3.1), Inches(1.6), Inches(6.6)],
        note="Business logic never lives in a route: churn_service.py holds plain functions taking a Session and returning a dict — the exact shape Phase 8 needed to turn endpoints into Claude tools with zero duplicated logic.",
    )

    # ------------------------------------------------- 6. PHASE 4 UI -----
    table_slide(
        prs, "Phase 4 — React Dashboard", "Five tabs, one shared fetch pattern, no router needed", "done",
        headers=["Tab", "Calls", "What it shows"],
        rows=[
            ("Churn Summary", "GET /churn/summary", "KPI cards + churn-rate bar per contract type"),
            ("Customer Search", "GET /customers/{id}", "Profile card, red/green by churn status, graceful not-found"),
            ("High-Risk Customers", "GET /customers/high-risk", "Sortable table, Load More paging, formatted currency"),
            ("Churn Prediction", "POST /predict-churn", "4-field form → colour-coded risk bar"),
            ("Retention Assistant", "POST /assistant/chat", "Chat UI — every reply shows which tools produced it"),
        ],
        col_widths=[Inches(2.9), Inches(3.2), Inches(5.2)],
        note="Built with Vite + React 19. Navigation is one activeTab state variable in App.jsx — deliberately no React Router.",
    )

    # ------------------------------------------------- 7. PHASE 5 DE -----
    table_slide(
        prs, "Phase 5 — Data Engineering", "A repeatable, quality-gated, now-orchestrated pipeline", "done",
        headers=["Stage", "What it proves"],
        rows=[
            ("Ingestion (dataengineer/ingestion.py)", "detect → validate schema (21 expected columns) → load → log"),
            ("Cleaning & curation (build_curated.py)", "Reuses CustomerCleaner, builds dim/fact tables + the high-risk view"),
            ("Feature build", "Same 6 engineered features from Phase 1, written to customer_ml_features"),
            ("PySpark cross-check (de5.ipynb)", "Same churn-rate aggregations on a second engine, written to Parquet"),
            ("Incremental load (de6.ipynb)", "Upsert logic — re-running doesn't duplicate customers"),
            ("Quality gate (build_curated.quality_report)", "Programmatic PASS/FAIL checks; raises to block downstream tasks"),
            ("Orchestration (dags/customer_pipeline_dag.py)", "Airflow DAG: ingest → validate → clean → score → daily brief → notify, @daily"),
        ],
        col_widths=[Inches(4.6), Inches(6.6)],
        note="The DAG wires in Phase 8's daily brief as its own scheduled task — pipeline and AI layer share one schedule.",
    )

    # ------------------------------------------------- 8. PHASE 6 ML -----
    bullets_slide(
        prs, "Phase 6 — Machine Learning", "A trained, explainable, batch-scored model", "done",
        items=[
            ("Problem framing: ", "drop customer_id (identifier, not signal), one-hot encode contract_type and "
             "internet_service, evaluate on F1 for the churned class (≈0.57–0.58) — never raw accuracy, "
             "since predicting “No” for everyone scores ~73% and is useless."),
            ("Two models trained & saved: ", "logistic_churn.pkl and tree_churn.pkl, plus feature_columns.json "
             "recording the exact training column order."),
            ("train.py: ", "runnable standalone from a cold terminal — persists the real feature list to "
             "models/feature_columns.json so predict.py never has to guess it."),
            ("predict.py: ", "loads the model once, builds a one-row feature vector, supports both an exact "
             "prediction (real customer_id) and a hypothetical what-if (four inputs + defaulted assumptions)."),
            ("batch_score.py: ", "applies predict_churn() to every row of customer_ml_features.csv, writes "
             "customer_risk_table.csv and dated snapshots to ml/risk_history/ for the daily brief to diff."),
        ],
    )

    # ------------------------------------------------- 9. PHASE 7 CL1 ----
    s = add_slide(prs)
    section_header(s, "Phase 7 — Lab CL1", "Claude audits the code — and finds a real bug", "done")
    bullet_block(s, Inches(0.9), Inches(2.0), Inches(6.7), Inches(4.6), [
        ("The defect: ", "predict.py only ever set 4 of the model's real inputs from the request — "
         "every other trained feature (total_charges, high_charge_flag, is_long_term_customer, "
         "has_streaming_bundle, auto_pay_flag, internet_service dummies) silently stayed at 0."),
        ("audit_defect.py: ", "pastes train.py + predict.py into one Claude request, forces a structured "
         "findings list via tool_choice — never free-text prose."),
        ("Scoped correctly: ", "findings explicitly mark this an inference-time bug — is_training_problem "
         "is forced false, so the fix is never “just retrain.”"),
        ("The fix: ", "predict.py now reads the real feature list from feature_columns.json and computes "
         "every engineered feature from the actual request, instead of guessing."),
        ("Proof, not a claim: ", "before/after predictions re-run and committed separately (docs/cl1_before_after.md)."),
    ])
    card_w2 = Inches(4.2)
    left2 = Inches(8.0)
    before = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left2, Inches(2.1), card_w2, Inches(1.9))
    before.adjustments[0] = 0.08; before.fill.solid(); before.fill.fore_color.rgb = SURFACE
    before.line.color.rgb = LINE; before.line.width = Pt(1); before.shadow.inherit = False
    tf = before.text_frame; tf.margin_left = Inches(0.2); tf.margin_top = Inches(0.15)
    p = tf.paragraphs[0]; r = p.add_run(); r.text = "BEFORE THE FIX"
    r.font.size = Pt(11); r.font.bold = True; r.font.name = MONO_FONT; r.font.color.rgb = INK_FAINT
    p2 = tf.add_paragraph(); r2 = p2.add_run(); r2.text = "45.8%"
    r2.font.size = Pt(34); r2.font.bold = True; r2.font.name = HEAD_FONT; r2.font.color.rgb = INK_SOFT
    p3 = tf.add_paragraph(); r3 = p3.add_run(); r3.text = "“Unlikely to churn”"
    r3.font.size = Pt(13); r3.font.name = BODY_FONT; r3.font.color.rgb = INK

    after = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left2, Inches(4.2), card_w2, Inches(1.9))
    after.adjustments[0] = 0.08; after.fill.solid(); after.fill.fore_color.rgb = SURFACE
    after.line.color.rgb = ACCENT; after.line.width = Pt(1.25); after.shadow.inherit = False
    tf = after.text_frame; tf.margin_left = Inches(0.2); tf.margin_top = Inches(0.15)
    p = tf.paragraphs[0]; r = p.add_run(); r.text = "AFTER THE FIX"
    r.font.size = Pt(11); r.font.bold = True; r.font.name = MONO_FONT; r.font.color.rgb = ACCENT
    p2 = tf.add_paragraph(); r2 = p2.add_run(); r2.text = "56.5%"
    r2.font.size = Pt(34); r2.font.bold = True; r2.font.name = HEAD_FONT; r2.font.color.rgb = ENERGY
    p3 = tf.add_paragraph(); r3 = p3.add_run(); r3.text = "“Likely to churn”"
    r3.font.size = Pt(13); r3.font.name = BODY_FONT; r3.font.color.rgb = INK
    add_text(s, "Same customer, same code path — the corrected feature vector changes the decision.",
              left2, Inches(6.25), card_w2, Inches(0.5), 11.5, INK_FAINT, italic=True, font=BODY_FONT)

    # ------------------------------------------------ 10. PHASE 7 CL2 ----
    bullets_slide(
        prs, "Phase 7 — Lab CL2", "Memory, versioned prompts, and a security checklist — in code", "done",
        items=[
            ("project_context.py: ", "under 200 lines of durable facts (dataset quirks, engineered features, "
             "table names, ground rules) prepended to every system prompt — no script re-explains the "
             "project from scratch, and none can drift from another's explanation."),
            ("prompts/*.txt + load_template(): ", "three versioned templates — profile the dataset, scaffold "
             "an endpoint, review code for security — prompt text is never inlined in a script."),
            ("Security checklist, enforced by code (security/guard.py):"),
            ("  1. ", "assert_no_secrets() — refuses to send a prompt containing a real .env secret value"),
            ("  2. ", "validate_tool_args() — every tool call's arguments checked against its own JSON Schema"),
            ("  3. ", "execute_tool() whitelists exactly which fields reach churn_service.py per tool"),
            ("  4. ", "cap_history() + config.ASSISTANT_MAX_* — one shared place for every token/turn limit"),
            ("Proven, not assumed: ", "the lab deliberately tries to leak a .env value through a template and "
             "confirms the guard blocks it."),
        ],
    )

    # ------------------------------------------------ 11. PHASE 8 AI1 ----
    bullets_slide(
        prs, "Phase 8 — Lab AI1", "The first Claude call from the backend", "done",
        items=[
            ("api/assistant.py: ", "one Anthropic() client built once at module level — not re-created per request."),
            ("Honest system prompt: ", "states the business facts plainly (telecom operator, ~7,043 customers, "
             "26.5% baseline churn) plus one hard rule — never invent customer data."),
            ("Model choice, documented: ", "claude-sonnet-5 for the tool-calling assistant and code review "
             "(ANTHROPIC_MODEL_STRONG); claude-haiku-4-5 as the cheaper default for lighter tasks."),
            ("Thinking on/off: ", "chat() accepts use_thinking, adding a 4096-token thinking budget when a "
             "harder question calls for visible reasoning."),
            ("Prompt caching, measured: ", "the system prompt is marked cache_control: ephemeral, and every "
             "response's cache-read / cache-write token counts are captured to compute a real cost-per-1,000-requests figure."),
        ],
    )

    # ------------------------------------------------ 12. PHASE 8 AI2 ----
    table_slide(
        prs, "Phase 8 — Lab AI2", "Claude's own APIs, given back to it as tools", "done",
        headers=["Tool", "What it calls"],
        rows=[
            ("get_customer_profile", "churn_service lookup by customer_id — reports “not found” honestly, never invents a profile"),
            ("get_churn_summary", "Company-wide KPIs — for any trend/rate/segment question, never estimated from memory"),
            ("get_high_risk_customers", "Rule-based v_high_risk_customers list, capped at a small limit, each with a risk_reason"),
            ("predict_churn", "Exact (real customer_id) or hypothetical (4 inputs) — hypothetical mode must surface its assumptions"),
        ],
        col_widths=[Inches(3.6), Inches(7.7)],
        note="One dispatch point, execute_tool(name, args): validates arguments against the tool's own schema, then calls straight into the SAME churn_service.py functions the React dashboard uses — zero duplicated business logic. The request loop calls Claude, executes every requested tool, sends results back, and repeats, capped at ASSISTANT_MAX_TOOL_ITERATIONS.",
    )

    # ------------------------------------------------ 13. PHASE 8 AI3 ----
    bullets_slide(
        prs, "Phase 8 — Lab AI3", "Shipped in the UI — with visible guardrails", "done",
        items=[
            ("A 5th dashboard tab: ", "AssistantChat.jsx — message list, input box, starter questions, retry-on-error."),
            ("Traceable answers: ", "every reply shows which tools produced it, so a support agent can trace any number back to a real query — the tool_calls trail returned by chat()."),
            ("Guardrails in the system prompt: ", "never state a number that didn't come from a tool call; cap any customer list at a small size; never speculate about a churn cause a tool result doesn't support."),
            ("Bounded conversation: ", "only the last 8 turns are sent with each request (MAX_HISTORY_TURNS in the UI, matched by config.ASSISTANT_MAX_HISTORY_TURNS on the backend) — cost and context stay predictable no matter how long the chat runs."),
            ("Stress-tested on purpose: ", "tricky asks like “list every customer” or “why did they churn” — a causal claim no tool can support — are exactly what Lab AI5's evaluation set checks."),
        ],
    )

    # ------------------------------------------------ 14. PHASE 8 AI4 ----
    s = add_slide(prs)
    section_header(s, "Phase 8 — Lab AI4", "Headless Claude: no chat UI at all", "done")
    left_w = Inches(5.9)
    code_card(s, Inches(0.9), Inches(2.0), left_w, Inches(4.6), [
        "pipeline/daily_brief.py",
        "",
        "1. Diff today's vs. yesterday's",
        "   risk_history snapshot",
        "2. Aggregate by contract segment",
        "   (avg risk, likely-to-churn count)",
        "3. Send ONLY the aggregated deltas",
        "   to Claude -- never raw rows",
        "4. tool_choice forces one fixed",
        "   shape: headline, movements,",
        "   recommended_actions",
        "",
        "Wired as its own task in the",
        "Airflow DAG, after batch scoring.",
    ], title="Daily Brief")
    code_card(s, Inches(6.95), Inches(2.0), left_w, Inches(4.6), [
        "ci/review_diff.py",
        "",
        "git diff HEAD~1 | python \\",
        "  ci/review_diff.py",
        "",
        "1. Reads a diff from stdin",
        "2. Forces structured findings via",
        "   tool_choice (file, issue,",
        "   severity) -- never free prose",
        "3. Empty diff -> no findings,",
        "   never invents an issue",
        "4. Exits non-zero on a real defect",
        "",
        "Wired into .githooks/pre-push --",
        "runs automatically before push.",
    ], title="CI Code Review")

    # ------------------------------------------------ 15. PHASE 8 AI5 ----
    bullets_slide(
        prs, "Phase 8 — Lab AI5", "Evaluated, cost-compared, and hardened", "done",
        items=[
            ("15 hand-written questions, ", "eval/questions.json — 5 single-tool, 5 multi-tool, 5 genuinely unanswerable."),
            ("Three scored rates: ", "factual accuracy (every number traces to a real tool result), correct tool "
             "selection, correct refusal on the unanswerable group."),
            ("--compare-models: ", "runs the full set on both claude-haiku-4-5 and claude-sonnet-5 to justify a "
             "production model choice on quality vs. cost, not guesswork."),
            ("Hardening, all in code: ", "response/history length limits, a tool-loop iteration cap, "
             "retry-with-backoff on rate-limit/overload errors, schema-validated tool arguments, a full audit "
             "log (audit_log.py — question, tools, tokens, per turn)."),
            ("docs/LIMITATIONS.md: ", "written honestly — what the assistant genuinely cannot know (why "
             "customers really leave, anything about the future or competitors) and what it should never be "
             "used for (any individual employment/credit/insurance decision)."),
        ],
    )

    # ------------------------------------------------ 16. SECURITY -------
    s = add_slide(prs)
    section_header(s, "Cross-Cutting", "Security & guardrails, enforced in code everywhere", "done")
    items = [
        ("No secrets in prompts", "assert_no_secrets() checks every outbound prompt against the live .env values (API keys, DB password) — not their names, their actual current values."),
        ("Schema-validated tool calls", "validate_tool_args() runs jsonschema against each tool's own input_schema before anything touches churn_service.py."),
        ("Bounded everything", "History turns, response tokens, tool iterations, and list sizes are each a single config.ASSISTANT_MAX_* constant — one place to change, everywhere enforced."),
        ("Full audit trail", "audit_log.py records every assistant turn: the question, every tool call and result, token usage, iteration count."),
        ("Config over hardcoding", "config.py loads DB credentials, the API key, and the Anthropic key from .env with no working fallback for secrets — anthropic_client() fails loudly if ANTHROPIC_API_KEY is missing."),
        ("Graceful tool failures", "execute_tool() never raises out to the caller — bad args, an unknown customer, or a real exception all come back as {\"error\": ...} so the loop can hand it to Claude instead of crashing."),
    ]
    card_w, card_h = Inches(3.85), Inches(2.35)
    gap_x, gap_y = Inches(0.2), Inches(0.2)
    left0 = (SLIDE_W - (card_w * 3 + gap_x * 2)) / 2
    top0 = Inches(1.95)
    for i, (h, d) in enumerate(items):
        col, row = i % 3, i // 3
        left = left0 + col * (card_w + gap_x)
        top = top0 + row * (card_h + gap_y)
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, card_w, card_h)
        card.adjustments[0] = 0.06; card.fill.solid(); card.fill.fore_color.rgb = SURFACE
        card.line.color.rgb = LINE; card.line.width = Pt(1); card.shadow.inherit = False
        tf = card.text_frame; tf.word_wrap = True
        tf.margin_left = Inches(0.18); tf.margin_right = Inches(0.16); tf.margin_top = Inches(0.15)
        p = tf.paragraphs[0]; r = p.add_run(); r.text = h
        r.font.size = Pt(14.5); r.font.bold = True; r.font.name = HEAD_FONT; r.font.color.rgb = ACCENT_STRONG
        p2 = tf.add_paragraph(); p2.space_before = Pt(5); r2 = p2.add_run(); r2.text = d
        r2.font.size = Pt(11.5); r2.font.name = BODY_FONT; r2.font.color.rgb = INK_SOFT

    # ------------------------------------------------ 17. ORCHESTRATION --
    s = add_slide(prs)
    section_header(s, "Orchestration", "One Airflow DAG, run @daily", "done")
    tasks = ["ingest", "validate_\nquality", "clean_and_\nbuild_curated", "build_features_\nand_score", "daily_\nbrief", "notify"]
    n = len(tasks)
    y = Inches(3.4)
    left0, right0 = Inches(1.1), SLIDE_W - Inches(1.1)
    step = (right0 - left0) / (n - 1)
    line = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, left0, y, right0, y)
    line.line.color.rgb = LINE; line.line.width = Pt(2)
    for i, label in enumerate(tasks):
        cx = left0 + step * i
        color = ENERGY if label.startswith("daily") else ACCENT
        dot = s.shapes.add_shape(MSO_SHAPE.OVAL, cx - Inches(0.1), y - Inches(0.1), Inches(0.2), Inches(0.2))
        dot.fill.solid(); dot.fill.fore_color.rgb = color; dot.line.fill.background(); dot.shadow.inherit = False
        add_text(s, label, cx - Inches(0.95), y + Inches(0.25), Inches(1.9), Inches(0.7),
                  13, INK, bold=(color == ENERGY), font=HEAD_FONT if color == ENERGY else BODY_FONT,
                  align=PP_ALIGN.CENTER)
    add_text(s, "dags/customer_pipeline_dag.py — correct, runnable Airflow code (not a mock). "
                 "The daily_brief task is Lab AI4's headless Claude call, scheduled right after batch scoring.",
              Inches(1.3), Inches(4.8), SLIDE_W - Inches(2.6), Inches(0.6), 13, INK_FAINT, italic=True,
              font=BODY_FONT, align=PP_ALIGN.CENTER)
    add_text(s, "quality_report() raises on failure — a bad load blocks every downstream task, it never silently continues.",
              Inches(1.3), Inches(5.45), SLIDE_W - Inches(2.6), Inches(0.6), 13, INK_FAINT, italic=True,
              font=BODY_FONT, align=PP_ALIGN.CENTER)

    # ------------------------------------------------ 18. TECH STACK -----
    s = add_slide(prs)
    section_header(s, "Under The Hood", "Technology stack")
    stack = [
        ("Language & core", "Python 3.12, pandas, SQLAlchemy 2.x"),
        ("Database", "MySQL (\"proj\" schema), pymysql driver"),
        ("API", "FastAPI, Pydantic v2 field validators"),
        ("Frontend", "React 19, Vite, plain fetch — no router"),
        ("ML", "scikit-learn (Logistic Regression + Decision Tree), joblib"),
        ("Data engineering", "Jupyter notebooks, PySpark, Parquet"),
        ("Orchestration", "Apache Airflow (DAG, @daily schedule)"),
        ("AI layer", "Anthropic Claude API — tool use, thinking, prompt caching"),
        ("Hardening", "jsonschema validation, custom retry/backoff, audit logging"),
    ]
    bullet_block(s, Inches(0.9), Inches(2.0), SLIDE_W - Inches(1.8), Inches(4.9),
                 [(f"{k}: ", v) for k, v in stack], size=15, gap=Pt(10))

    # ------------------------------------------------ 19. LIMITATIONS ----
    bullets_slide(
        prs, "Honest Boundaries", "What this system does not claim to do", "warn",
        items=[
            ("No causal explanations: ", "there's no exit-interview or free-text feedback data — the "
             "assistant can state correlations a tool already encodes, never a true reason a specific customer left."),
            ("No forecasting: ", "“what will churn look like next quarter” has no tool that can answer it "
             "truthfully — the assistant is instructed to refuse rather than extrapolate."),
            ("Single shared API key: ", "validate_tool_args() checks argument shape, not caller authorization — "
             "there's no per-user permission model yet."),
            ("Orchestration is real but requires Airflow: ", "the DAG is correct, runnable code, but exercising "
             "it for real means an actual Airflow install (`airflow db init`) pointed at this repo."),
            ("Evaluation is a grounding proxy: ", "“factual accuracy” checks that numbers trace back to a "
             "real tool result, not full semantic correctness of the sentence around them."),
        ],
    )

    # ------------------------------------------------ 20. CLOSE ----------
    s = add_slide(prs)
    add_text(s, "One service layer.", Inches(1.3), Inches(2.3), SLIDE_W - Inches(2.6), Inches(0.9),
              34, INK, bold=True, font=HEAD_FONT, align=PP_ALIGN.CENTER)
    add_text(s, "Every tab, every tool, every scheduled job calls the same churn_service.py functions.",
              Inches(1.8), Inches(3.2), SLIDE_W - Inches(3.6), Inches(0.7), 17, INK_SOFT,
              font=BODY_FONT, align=PP_ALIGN.CENTER)
    hr(s, Inches(4.5), Inches(4.15), Inches(4.333))
    add_text(s, "That discipline, paid for in Phase 3, is what made Phases 7 and 8 a schema-wrapping "
                 "exercise — not a rewrite.",
              Inches(2.3), Inches(4.4), SLIDE_W - Inches(4.6), Inches(0.9), 15, ACCENT_STRONG,
              font=BODY_FONT, align=PP_ALIGN.CENTER, italic=True)

    prs.save(OUT_PATH)
    print(f"Wrote {OUT_PATH} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    build()
