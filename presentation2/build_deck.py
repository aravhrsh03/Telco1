"""
Build the sales/pitch deck for presentation2 - "Network Operations &
Predictive Intelligence", feature by feature.

    python presentation2/build_deck.py
    -> presentation2/Network_Operations_Predictive_Intelligence.pptx

Content mirrors PITCH_DECK.md / SCRIPT.md in this folder. All slide text
lives in the function bodies below and in the BUILD section at the bottom -
edit there, then re-run.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# --------------------------------------------------------------------------
# palette
# --------------------------------------------------------------------------
NAVY = RGBColor(0x11, 0x2A, 0x43)
STEEL = RGBColor(0x2E, 0x6C, 0x8E)
AMBER = RGBColor(0xDD, 0x7A, 0x2E)
LIGHT = RGBColor(0xEE, 0xF2, 0xF5)
GRAYTX = RGBColor(0x53, 0x62, 0x6E)
DARKTX = RGBColor(0x1E, 0x28, 0x30)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RULE = RGBColor(0xDD, 0x7A, 0x2E)

FONT = "Segoe UI"
FONT_LIGHT = "Segoe UI Light"

EMU_W = Inches(13.333)
EMU_H = Inches(7.5)

prs = Presentation()
prs.slide_width = EMU_W
prs.slide_height = EMU_H
BLANK = prs.slide_layouts[6]


# --------------------------------------------------------------------------
# helpers
# --------------------------------------------------------------------------
def _set(run, size, *, bold=False, color=DARKTX, font=FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def rect(slide, x, y, w, h, fill, line=None, line_w=None, shape=MSO_SHAPE.RECTANGLE):
    sp = slide.shapes.add_shape(shape, x, y, w, h)
    sp.fill.solid()
    sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = line_w or Pt(1)
    sp.shadow.inherit = False
    return sp


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    return tf


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text.strip()


def base_slide(title, kicker="", *, number=None):
    slide = prs.slides.add_slide(BLANK)
    rect(slide, 0, 0, EMU_W, Inches(1.12), NAVY)
    rect(slide, 0, Inches(1.12), EMU_W, Pt(3), RULE)
    tf = textbox(slide, Inches(0.6), Inches(0.14), Inches(12.1), Inches(0.9),
                 anchor=MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    if kicker:
        r = p.add_run(); _set(r, 12, bold=True, color=AMBER, font=FONT)
        r.text = kicker.upper() + "   "
    r = p.add_run(); _set(r, 26, bold=True, color=WHITE, font=FONT)
    r.text = title
    ft = textbox(slide, Inches(0.6), Inches(7.02), Inches(9), Inches(0.4))
    fr = ft.paragraphs[0].add_run()
    _set(fr, 9, color=GRAYTX, font=FONT)
    fr.text = "Network Operations & Predictive Intelligence"
    if number is not None:
        nt = textbox(slide, Inches(12.4), Inches(7.02), Inches(0.6), Inches(0.4))
        nt.paragraphs[0].alignment = PP_ALIGN.RIGHT
        nr = nt.paragraphs[0].add_run()
        _set(nr, 9, color=GRAYTX, font=FONT)
        nr.text = str(number)
    return slide


def key_message(slide, text, *, y=Inches(1.42)):
    tf = textbox(slide, Inches(0.6), y, Inches(12.1), Inches(0.8))
    r = tf.paragraphs[0].add_run()
    _set(r, 18, bold=True, color=STEEL, font=FONT)
    r.text = text
    return tf


def bullets(slide, items, *, x=Inches(0.6), y=Inches(2.35),
            w=Inches(8.2), h=Inches(4.2), size=15, gap=10):
    tf = textbox(slide, x, y, w, h)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        if isinstance(item, tuple):
            text, lvl = item
        else:
            text, lvl = item, 0
        p.level = lvl
        dot = p.add_run()
        _set(dot, size, bold=True, color=AMBER if lvl == 0 else STEEL)
        dot.text = ("•  " if lvl == 0 else "–  ")
        r = p.add_run()
        _set(r, size, color=DARKTX if lvl == 0 else GRAYTX)
        r.text = text
    return tf


def callout(slide, lines, *, x=Inches(9.2), y=Inches(2.35),
            w=Inches(3.55), h=Inches(2.1), title=None):
    box = rect(slide, x, y, w, h, LIGHT, line=STEEL, line_w=Pt(1),
               shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.22)
    tf.margin_right = Inches(0.18)
    tf.margin_top = Inches(0.18)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    first = True
    if title:
        p = tf.paragraphs[0]; first = False
        r = p.add_run(); _set(r, 11, bold=True, color=STEEL); r.text = title.upper()
        p.space_after = Pt(8)
    for ln in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        r = p.add_run(); _set(r, 13, color=DARKTX); r.text = ln
        p.space_after = Pt(6)
    return box


def metric_tiles(slide, tiles, *, y=Inches(2.5), h=Inches(1.9)):
    n = len(tiles)
    gap = Inches(0.3)
    total = EMU_W - Inches(1.2)
    w = Emu(int((total - gap * (n - 1)) / n))
    x = Inches(0.6)
    for value, label in tiles:
        box = rect(slide, x, y, w, h, NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); _set(r, 32, bold=True, color=WHITE); r.text = value
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
        p2.space_before = Pt(4)
        r2 = p2.add_run(); _set(r2, 11.5, color=RGBColor(0xC5, 0xD3, 0xDE))
        r2.text = label
        x = Emu(x + w + gap)


def two_col_table(slide, rows, *, y=Inches(2.4), row_h=Inches(0.66)):
    x1, w1 = Inches(0.6), Inches(4.6)
    x2, w2 = Inches(5.35), Inches(7.35)
    cy = y
    for left, right in rows:
        b1 = rect(slide, x1, cy, w1, row_h, NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        tf1 = b1.text_frame; tf1.vertical_anchor = MSO_ANCHOR.MIDDLE; tf1.word_wrap = True
        p1 = tf1.paragraphs[0]; p1.alignment = PP_ALIGN.CENTER
        r1 = p1.add_run(); _set(r1, 14, bold=True, color=WHITE); r1.text = left
        b2 = rect(slide, x2, cy, w2, row_h, LIGHT, line=STEEL, line_w=Pt(0.75),
                  shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        tf2 = b2.text_frame; tf2.vertical_anchor = MSO_ANCHOR.MIDDLE; tf2.word_wrap = True
        tf2.margin_left = Inches(0.2); tf2.margin_right = Inches(0.2)
        p2 = tf2.paragraphs[0]
        r2 = p2.add_run(); _set(r2, 13, color=DARKTX); r2.text = right
        cy = Emu(cy + row_h + Inches(0.12))


def feature_slide(n, number_label, title, what_it_does, why_it_matters, proof=None):
    """Standard 'Feature N' slide: What it does / Why it matters / optional proof line."""
    s = base_slide(title, f"Feature {number_label}", number=n)
    y = Inches(1.55)
    head1 = textbox(s, Inches(0.6), y, Inches(12.1), Inches(0.4))
    hr = head1.paragraphs[0].add_run(); _set(hr, 13, bold=True, color=AMBER)
    hr.text = "WHAT IT DOES"
    tf1 = textbox(s, Inches(0.6), Emu(y + Inches(0.42)), Inches(12.1), Inches(1.4))
    r1 = tf1.paragraphs[0].add_run(); _set(r1, 15.5, color=DARKTX); r1.text = what_it_does

    y2 = Emu(y + Inches(2.0))
    head2 = textbox(s, Inches(0.6), y2, Inches(12.1), Inches(0.4))
    hr2 = head2.paragraphs[0].add_run(); _set(hr2, 13, bold=True, color=STEEL)
    hr2.text = "WHY IT MATTERS"
    tf2 = textbox(s, Inches(0.6), Emu(y2 + Inches(0.42)), Inches(12.1), Inches(1.6))
    r2 = tf2.paragraphs[0].add_run(); _set(r2, 15.5, color=DARKTX); r2.text = why_it_matters

    if proof:
        box = rect(s, Inches(0.6), Inches(6.05), Inches(12.1), Inches(0.75), LIGHT,
                   line=STEEL, line_w=Pt(1), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        tf3 = box.text_frame; tf3.vertical_anchor = MSO_ANCHOR.MIDDLE; tf3.word_wrap = True
        tf3.margin_left = Inches(0.2); tf3.margin_right = Inches(0.2)
        p = tf3.paragraphs[0]
        d = p.add_run(); _set(d, 12.5, bold=True, color=STEEL); d.text = "PROOF POINT   "
        r = p.add_run(); _set(r, 12.5, color=DARKTX); r.text = proof
    return s


def live_demo_slide(n, title, cues):
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, EMU_W, EMU_H, NAVY)
    rect(s, 0, Inches(1.5), EMU_W, Pt(3), RULE)
    tf = textbox(s, Inches(0.7), Inches(0.35), Inches(10), Inches(1.0),
                 anchor=MSO_ANCHOR.MIDDLE)
    r = tf.paragraphs[0].add_run()
    _set(r, 13, bold=True, color=AMBER)
    r.text = "LIVE DEMO"
    tf2 = textbox(s, Inches(0.7), Inches(0.7), Inches(11.5), Inches(0.8))
    r2 = tf2.paragraphs[0].add_run()
    _set(r2, 26, bold=True, color=WHITE, font=FONT_LIGHT)
    r2.text = title
    body = textbox(s, Inches(0.9), Inches(2.1), Inches(11.4), Inches(4.6))
    for i, cue in enumerate(cues):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.space_after = Pt(12)
        d = p.add_run(); _set(d, 15, bold=True, color=AMBER); d.text = f"{i + 1}.  "
        r3 = p.add_run(); _set(r3, 15, color=RGBColor(0xE7, 0xEC, 0xF0)); r3.text = cue
    nt = textbox(s, Inches(12.4), Inches(7.02), Inches(0.6), Inches(0.4))
    nt.paragraphs[0].alignment = PP_ALIGN.RIGHT
    nr = nt.paragraphs[0].add_run()
    _set(nr, 9, color=RGBColor(0x8F, 0xA6, 0xB6), font=FONT)
    nr.text = str(n)
    return s


# --------------------------------------------------------------------------
# TITLE SLIDE
# --------------------------------------------------------------------------
def title_slide():
    slide = prs.slides.add_slide(BLANK)
    rect(slide, 0, 0, EMU_W, EMU_H, NAVY)
    rect(slide, 0, Inches(4.35), EMU_W, Pt(3), RULE)
    tf = textbox(slide, Inches(0.9), Inches(2.1), Inches(11.5), Inches(2.2))
    r = tf.paragraphs[0].add_run()
    _set(r, 44, bold=True, color=WHITE, font=FONT_LIGHT)
    r.text = "Network Operations &\nPredictive Intelligence"
    sub = textbox(slide, Inches(0.95), Inches(4.6), Inches(11.4), Inches(1.4))
    sr = sub.paragraphs[0].add_run()
    _set(sr, 17, color=RGBColor(0xC5, 0xD3, 0xDE))
    sr.text = ("A live NOC platform that doesn't just show you the network - "
               "it tells you what's happening, whether it's normal, and what "
               "to check next.")
    meta = textbox(slide, Inches(0.95), Inches(6.2), Inches(11.4), Inches(0.6))
    mr = meta.paragraphs[0].add_run()
    _set(mr, 12, color=RGBColor(0x8F, 0xA6, 0xB6))
    mr.text = "Presented by [Your Name]  ·  [Date]"
    add_notes(slide, """
"Hi, I'm [Your Name]. This is Network Operations and Predictive Intelligence -
a platform that doesn't just show you the network, it tells you what's
happening, whether it's normal, and what to check next. Let me show you,
before I explain how it works."
""")


def slide_hook(n):
    live_demo_slide(n, "The Hook (~75 sec)", [
        "Switch to the live dashboard - open Hotspots & Milan Map.",
        "Point at the map: \"That's a real 10,000-cell grid over Milan, ranked by real activity, right now.\"",
        "Click the top-ranked hotspot polygon.",
        "Switch to the AI assistant window and ask: \"Explain grid [chosen id] right now.\"",
        "Let the four-part answer render - don't summarize it, let it speak for itself.",
        "Cut back to the deck: \"Every number in that came from the exact platform you just saw.\"",
    ])
    add_notes(prs.slides[-1], """
Don't explain anything yet - just show it. Then: "That explanation didn't come
from a script. Every number in it was pulled live from the exact platform you
just saw. Now let me walk you through it, feature by feature."
""")


def slide_problem(n):
    s = base_slide("The Problem", "Why this exists", number=n)
    key_message(s, "Every NOC team deals with the same three gaps.")
    bullets(s, [
        "Raw numbers, no context - a dashboard can show a spike, but not whether it's normal for that cell at that hour.",
        "No trust signal - if the pipeline had a bad run today, most dashboards show yesterday's numbers with total confidence and no warning.",
        "Manual investigation - an engineer opens five tools, pulls history, checks a model score, and writes it up, every single time.",
    ], y=Inches(2.5), w=Inches(12.0))
    tf = textbox(s, Inches(0.6), Inches(5.6), Inches(12.1), Inches(1.1))
    r = tf.paragraphs[0].add_run()
    _set(r, 15, bold=True, color=STEEL)
    r.text = ("What a NOC team actually needs: a system that separates what "
              "happened from is this unusual from what does it mean - and never "
              "blurs the three.")
    add_notes(s, """
"Here's why that matters. Every NOC deals with the same problem: raw numbers
everywhere, but no easy way to know if a number is actually normal. A
dashboard can show you a spike - it can't tell you whether you can even trust
the data behind it, and it can't tell you what to do about it."
""")


def slide_promise(n):
    s = base_slide("The Promise", "One design boundary", number=n)
    key_message(s, "Every layer answers exactly one kind of question.")
    two_col_table(s, [
        ("Data + processing", "What happened?"),
        ("Machine learning", "Is this unusual or risky?"),
        ("AI assistant", "What does it mean? What should I check next?"),
    ], y=Inches(2.5))
    tf = textbox(s, Inches(0.6), Inches(5.1), Inches(12.1), Inches(1.4))
    r = tf.paragraphs[0].add_run()
    _set(r, 15, bold=True, color=STEEL)
    r.text = ('"Data tells you what happened. Machine learning tells you if '
              'it\'s unusual. The AI layer tells you what it means and what to '
              'check next - and it never does another layer\'s job."')
    add_notes(s, """
"So the whole product is built around one boundary. Data and processing tell
you what happened. Machine learning tells you if it's unusual or risky. And
the AI layer tells you what it means and what to check next. No layer does
another layer's job. That discipline is what you just watched work."
""")


def slide_feature1(n):
    feature_slide(n, "1", "Trusted Ingestion",
        "Every incoming daily activity file is validated before it can touch "
        "anything downstream - schema check, minimum-quality check (no malformed "
        "timestamps, no out-of-range grid IDs, no negative activity, no "
        "duplicate files). Anything that fails is quarantined with a named "
        "reason, never silently dropped and never silently processed.",
        "A bad file corrupting the analytics layer is the single most "
        "expensive kind of failure - because it doesn't announce itself. This "
        "feature exists so it can't happen quietly.",
        proof="Every file, good or bad, gets an audit row - filename, status, "
              "row count, reason, timestamp. Nothing enters the system unaccounted for.")
    add_notes(prs.slides[-1], """
"Let's go feature by feature, starting at the very front door. Every file
that comes in is validated before it touches anything downstream. Anything
bad is quarantined with a named reason, never silently dropped, and never
silently processed either."
""")


def slide_feature2(n):
    feature_slide(n, "2", "One Governed Pipeline, One Trust Signal",
        "Ingested files are cleaned and aggregated at real scale, loaded into "
        "a proper analytics warehouse, and checked by 15 automated quality "
        "gates after every run - grain correctness, row-count consistency, "
        "geographic sanity, cross-layer agreement. Any single gate failing "
        "fails the whole run.",
        "That produces one machine-readable artifact everything else depends "
        "on: a pipeline status record - \"is the data behind this dashboard "
        "trustworthy, right now?\" Not a log line - a first-class output the "
        "API, the dashboard, and the AI assistant all read before saying anything.",
        proof="The platform has processed 1.68 million grid-hour records end "
              "to end, guarded the whole way by those 15 gates.")
    add_notes(prs.slides[-1], """
"After ingestion, the data is processed at real scale and loaded into an
analytics warehouse. Then fifteen automated quality gates check the output -
any single failure fails the whole run. That produces a pipeline status
record - a machine-readable trust signal that the API, the dashboard, and the
AI assistant all check before they say anything."
""")


def slide_feature3(n):
    feature_slide(n, "3", "Live Network Summary",
        "One glance answers \"how's the network right now\" - total activity, "
        "how many grid cells are active, the peak hour, and the top cell for "
        "the current reporting window.",
        "This is the front door. It's the first thing an operator sees, and it "
        "always shows the data's own current time (not the viewer's clock) - "
        "critical for a system where \"now\" is a defined concept, not just "
        "whatever the browser says.")
    add_notes(prs.slides[-1], """
"That's what feeds the front page of the dashboard - total activity, how many
cells are active, the peak hour, the top cell, all for right now. Notice it
shows the data's current time, not the viewer's clock."
""")


def slide_feature4(n):
    feature_slide(n, "4", "Grid Explorer",
        "Pick any one of the 10,000 geographic grid cells covering Milan and "
        "pull its hourly activity - SMS, call, and internet activity, "
        "individually and combined - on demand.",
        "Every summary and every alert eventually leads an operator to ask "
        "\"wait, show me that specific cell.\" This is that drill-down, and "
        "it's just a thin, fast consumer of the same stable API everything "
        "else uses.")
    add_notes(prs.slides[-1], """
"Next, drill-down. Pick any one of the ten thousand grid cells and pull its
hourly activity - SMS, call, internet, and combined - on demand."
""")


def slide_feature5(n):
    feature_slide(n, "5", "Hotspot Ranking + the Milan Map",
        "Ranks grid cells by current activity and renders the highlighted "
        "cells on a real, interactive map of Milan's 100x100 grid.",
        "Each cell in the underlying map data carries two possible "
        "identifiers, and the wrong one would silently draw every hotspot on "
        "its neighbor's square - full apparent coverage, no error, no "
        "warning. This platform joins on the correct identifier and proves "
        "it: any highlighted cell's centerpoint can be pulled independently "
        "and checked against where it should actually sit.",
        proof="That's the difference between a map that looks right and a "
              "map that is right.")
    add_notes(prs.slides[-1], """
"Here's the detail worth dwelling on. Each cell in the underlying map data
carries two possible identifiers, and using the wrong one would silently draw
every hotspot on its neighbor's square - full apparent coverage, no error,
no warning anywhere. This platform joins on the correct one, and proves it."
""")


def slide_feature6(n):
    feature_slide(n, "6", "Rule-Based Alerts",
        "A transparent, explainable first layer of detection - each cell's "
        "current activity is compared against its own recent baseline, and "
        "three plain rules fire when activity is materially high, materially "
        "low, or spiking sharply versus the previous hour. Every alert "
        "carries a human-readable reason.",
        "No black box. An operator (or an auditor) can read the rule and the "
        "number and understand exactly why an alert fired - before any "
        "machine learning is involved at all.")
    add_notes(prs.slides[-1], """
"Before any machine learning gets involved, there's a transparent first
layer: each cell's current activity compared to its own recent baseline,
three plain rules, every alert carrying a human-readable reason."
""")


def slide_feature7(n):
    s = feature_slide(n, "7", "Predictive Risk Scoring",
        "A live machine-learning model estimates the probability that a grid "
        "cell is about to see an activity surge in the next hour, from six "
        "engineered features describing its trailing 24-hour behavior - how "
        "active, how peaky, how erratic, how fast-growing, how internet-heavy.",
        "The model predicts the next hour using only past data - never the "
        "same window it's being scored on - and it's evaluated on a strictly "
        "time-ordered split. That discipline is what makes the score real "
        "instead of a number that looks good on a slide and means nothing in "
        "production.")
    metric_tiles(s, [
        ("70.6%", "Test accuracy - deliberately below the 91% \"always quiet\" baseline"),
        ("68%", "Recall - catches ~2 in 3 real activity surges"),
        ("19.5%", "Precision - an attention filter, not an automated trigger"),
    ], y=Inches(6.05), h=Inches(0.85))
    add_notes(s, """
"Here's the number I actually want to talk about: seventy point six percent
test accuracy - on paper, worse than just predicting 'nothing's wrong' every
time, which scores about ninety percent. That's deliberate - tuned to catch
real surges at the cost of some false alarms. And if this number had come
back above ninety-five percent, that would mean the model was seeing its own
answer, and I'd have gone back to fix it."
""")


def slide_feature8_9(n):
    s = feature_slide(n, "8", "Anomaly Detection",
        "A second, independent signal - not \"is this predicted to surge,\" "
        "but \"is this unusual for this specific cell at this specific hour "
        "of day,\" based on its own history. A cell that's normally quiet at "
        "3am is judged against its own 3am baseline, not a network-wide average.",
        "Three signals - the transparent rule, the predictive model, and this "
        "historical-anomaly check - see three different things. When they "
        "agree, an operator has high confidence fast. When they disagree, "
        "that disagreement is surfaced as information, not smoothed away.")
    add_notes(s, """
"One more signal sits alongside the risk score: an anomaly baseline that asks
'is this unusual for this cell at this hour of day,' based on its own
history. Three independent signals - when they agree, that's high confidence
fast. When they disagree, that's information surfaced to a human, not
smoothed over."
""")


def slide_demo2(n):
    live_demo_slide(n, "Feature 9: The AI Assistant (~90-120 sec)", [
        "Switch to the AI assistant window.",
        "If pipeline status is healthy, ask: \"What's the network situation for the last hour?\"",
        "Narrate as it runs: pipeline status first, then summary, then hotspots - before it says a word.",
        "Point to the four sections as they render: Severity, Evidence, Interpretation, Next Checks.",
        "If pipeline status is unhealthy instead, ask the same question and narrate: \"Watch what happens when the pipeline itself isn't healthy\" - it reports the gap rather than answering anyway.",
        "Return to slides: \"Nothing here was invented - every number came from a real tool call.\"",
    ])
    add_notes(prs.slides[-1], """
"It checks the pipeline's trust signal first, then pulls current activity,
history, the model's risk score, the anomaly signal, and the cell's real
location - all through defined tool calls - and returns a structured answer.
It never sees raw data, never invents a number, and never asserts a network
situation as fact without checking whether the data can even be trusted."
""")


def slide_trust(n):
    s = base_slide("Why You Can Trust Every Number On Screen", "Rigor", number=n)
    bullets(s, [
        "The core grain rule - one row per grid cell per hour - is enforced in code, in three separate places, including the warehouse's own primary key.",
        "The risk model trains on a strictly time-ordered split; it never sees the future.",
        "The vocabulary is deliberately strict: this system never claims “congestion”, because there is no capacity or throughput data anywhere in it - only proportional activity measures.",
        "The pipeline has been tested against deliberately injected failures - a missing file, a duplicate, a malformed timestamp, negative values, a schema change - and each one produces a distinguishable, honest status.",
    ], y=Inches(2.5), w=Inches(12.0))
    add_notes(s, """
"A few things make this more than a demo. The one-row-per-cell-per-hour rule
is enforced in code, in three separate places, including the database's own
primary key. The risk model trains on a strictly time-ordered split. The
vocabulary is deliberately strict - this system never claims 'congestion.'
And the pipeline has actually been tested against injected failures."
""")


def slide_scale(n):
    s = base_slide("The Scale, Recap", "By the numbers", number=n)
    metric_tiles(s, [
        ("10,000", "Grid cells modeled (100x100 over Milan)"),
        ("1.68M", "Analytics records processed (grid-hours)"),
        ("15", "Quality gates per run - any failure fails the run"),
    ], y=Inches(2.4), h=Inches(1.85))
    metric_tiles(s, [
        ("6", "Stable API endpoints"),
        ("6", "ML features per prediction"),
        ("3", "Independent detection signals"),
    ], y=Inches(4.5), h=Inches(1.85))
    add_notes(s, """
"So, the numbers behind everything you just watched: ten thousand grid cells,
one point six eight million analytics records, fifteen automated quality
gates, six stable API endpoints, six engineered features, and three
independent signals working together."
""")


def slide_roadmap(n):
    s = base_slide("Where This Goes Next", "Roadmap", number=n)
    bullets(s, [
        "Wire the ML risk scores directly into the hotspot ranking view.",
        "Replace the remaining hardcoded alert thresholds with the full rule engine.",
        "Expose the platform's tools over a standard protocol (MCP) so any AI surface can investigate a grid the same, grounded way.",
        "Keep accumulating history - the anomaly baseline gets sharper with every additional day of data.",
        "Point ingestion at a real live feed - the architecture doesn't change, only the source.",
    ], y=Inches(2.5), w=Inches(12.0))
    add_notes(s, """
"And this isn't a finished, static thing. Next: wiring the ML risk scores
directly into the hotspot view, replacing the remaining hardcoded thresholds
with the full rule engine, exposing these tools over a standard protocol so
any AI surface can investigate a cell the same way, and simply accumulating
more history so the anomaly baseline keeps getting sharper."
""")


def slide_close(n):
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, EMU_W, EMU_H, NAVY)
    rect(s, Inches(0.9), Inches(2.5), Inches(0.12), Inches(2.4), RULE)
    tf = textbox(s, Inches(1.3), Inches(2.35), Inches(11.0), Inches(2.8),
                 anchor=MSO_ANCHOR.MIDDLE)
    r = tf.paragraphs[0].add_run()
    _set(r, 24, bold=True, color=WHITE, font=FONT_LIGHT)
    r.text = ("Data tells you what happened.\n"
              "Machine learning tells you if it's unusual.\n"
              "The AI layer tells you what it means and what to check next.")
    sub = textbox(s, Inches(1.3), Inches(5.2), Inches(11.0), Inches(0.8))
    sr = sub.paragraphs[0].add_run()
    _set(sr, 15, color=RGBColor(0xC5, 0xD3, 0xDE))
    sr.text = "Network Operations & Predictive Intelligence - without ever doing the analytics stack's job."
    q = textbox(s, Inches(1.3), Inches(6.2), Inches(11), Inches(0.6))
    qr = q.paragraphs[0].add_run()
    _set(qr, 13, bold=True, color=AMBER)
    qr.text = "Thank you - questions?"
    add_notes(s, """
"That's Network Operations and Predictive Intelligence - data tells you what
happened, machine learning tells you if it's unusual, and the AI layer tells
you what it means and what to check next, without ever doing the analytics
stack's job. Thanks for watching - happy to answer questions."
""")


# --------------------------------------------------------------------------
# BUILD
# --------------------------------------------------------------------------
title_slide()                  # 1
slide_hook(2)                  # 2  LIVE DEMO
slide_problem(3)                # 3
slide_promise(4)                # 4
slide_feature1(5)               # 5
slide_feature2(6)               # 6
slide_feature3(7)               # 7
slide_feature4(8)               # 8
slide_feature5(9)               # 9
slide_feature6(10)              # 10
slide_feature7(11)              # 11
slide_feature8_9(12)            # 12
slide_demo2(13)                 # 13  LIVE DEMO
slide_trust(14)                 # 14
slide_scale(15)                 # 15
slide_roadmap(16)               # 16
slide_close(17)                 # 17

OUT = "presentation2/Network_Operations_Predictive_Intelligence.pptx"
prs.save(OUT)
print(f"wrote {OUT}  ({len(prs.slides._sldIdLst)} slides)")
